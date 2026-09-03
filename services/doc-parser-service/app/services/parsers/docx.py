"""DOCX Parser using python-docx"""
from docx import Document
from docx.document import Document as _Document
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph
from loguru import logger
from typing import Tuple
import os


def _iter_block_items(parent):
    """Yield paragraphs and tables in their original DOCX document order."""
    if isinstance(parent, _Document):
        parent_elm = parent.element.body
    elif isinstance(parent, _Cell):
        parent_elm = parent._tc
    else:
        raise ValueError(f"unsupported DOCX parent: {type(parent)!r}")

    for child in parent_elm.iterchildren():
        if child.tag.endswith("}p"):
            yield Paragraph(child, parent)
        elif child.tag.endswith("}tbl"):
            yield Table(child, parent)


def _table_text(table: Table) -> str:
    """Render a Word table as tab-separated rows for retrieval and citations."""
    rows = []
    for row in table.rows:
        cells = []
        for cell in row.cells:
            # A cell may contain multiple paragraphs; keep them readable while
            # retaining the table's row/column boundaries.
            value = "\n".join(p.text.strip() for p in cell.paragraphs if p.text.strip())
            cells.append(value.replace("\t", " ").replace("\n", " / "))
        if any(cells):
            rows.append("\t".join(cells))
    return "\n".join(rows)


def parse_docx(file_path: str) -> Tuple[str, int]:
    """
    Parse DOCX file and extract text with structure
    
    Args:
        file_path: Path to DOCX file
        
    Returns:
        Tuple of (extracted_text, file_size_bytes)
    """
    try:
        doc = Document(file_path)
        text_parts = []
        
        for block in _iter_block_items(doc):
            if isinstance(block, Table):
                table_text = _table_text(block)
                if table_text:
                    text_parts.append(table_text)
                continue

            # Preserve heading structure
            if block.style.name.startswith('Heading'):
                level = block.style.name.replace('Heading ', '')
                prefix = '#' * int(level) if level.isdigit() else '#'
                text_parts.append(f"\n{prefix} {block.text}\n")
            else:
                text_parts.append(block.text)
        
        full_text = "\n".join(text_parts)
        file_size = os.path.getsize(file_path)
        
        logger.info(
            f"Parsed DOCX: {file_path}, paragraphs={len(doc.paragraphs)}, "
            f"tables={len(doc.tables)}, size={file_size} bytes"
        )
        return full_text, file_size
        
    except Exception as e:
        logger.error(f"Failed to parse DOCX {file_path}: {e}")
        raise
