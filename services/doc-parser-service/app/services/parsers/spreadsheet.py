"""Structured parsers for modern and legacy Excel workbooks."""

from datetime import date, datetime, time
from pathlib import Path
import shutil
import subprocess
import tempfile
from typing import Iterable, Sequence, Tuple

import xlrd
from loguru import logger
from openpyxl import load_workbook


MAX_ROWS_PER_SHEET = 100_000
MAX_COLUMNS_PER_SHEET = 1_024
MAX_EXTRACTED_CHARS = 10_000_000
CONVERSION_TIMEOUT_SECONDS = 60


class SpreadsheetConversionError(RuntimeError):
    """Raised when a compatibility-encrypted XLS cannot be converted safely."""


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip().replace("\r\n", "\n").replace("\r", "\n")


def _render_sheet(name: str, rows: Iterable[Sequence[object]]) -> str:
    lines = [f"--- Sheet: {name} ---"]
    extracted_chars = len(lines[0])
    for row_number, row in enumerate(rows, start=1):
        if row_number > MAX_ROWS_PER_SHEET:
            raise ValueError(f"sheet {name!r} exceeds {MAX_ROWS_PER_SHEET} rows")
        if len(row) > MAX_COLUMNS_PER_SHEET:
            raise ValueError(f"sheet {name!r} exceeds {MAX_COLUMNS_PER_SHEET} columns")
        cells = [_cell_text(value) for value in row]
        while cells and not cells[-1]:
            cells.pop()
        if not cells or not any(cells):
            continue
        line = "\t".join(cells)
        extracted_chars += len(line)
        if extracted_chars > MAX_EXTRACTED_CHARS:
            raise ValueError(f"sheet {name!r} exceeds extracted text limit")
        lines.append(line)
    return "\n".join(lines)


def _parse_xlsx(file_path: str) -> str:
    workbook = load_workbook(file_path, read_only=True, data_only=True)
    try:
        return "\n\n".join(
            _render_sheet(sheet.title, sheet.iter_rows(values_only=True))
            for sheet in workbook.worksheets
        )
    finally:
        workbook.close()


def _parse_xls(file_path: str) -> str:
    workbook = xlrd.open_workbook(file_path, on_demand=True)
    try:
        rendered = []
        for sheet in workbook.sheets():
            rows = (sheet.row_values(index) for index in range(sheet.nrows))
            rendered.append(_render_sheet(sheet.name, rows))
        return "\n\n".join(rendered)
    finally:
        workbook.release_resources()


def _parse_compatibility_encrypted_xls(file_path: str) -> str:
    libreoffice = shutil.which("libreoffice")
    if libreoffice is None:
        raise SpreadsheetConversionError("LibreOffice is required for compatibility-encrypted XLS")
    with tempfile.TemporaryDirectory(prefix="legacy-xls-") as temp_dir:
        work_dir = Path(temp_dir)
        source_path = work_dir / "source.xls"
        output_path = work_dir / "source.xlsx"
        profile_dir = work_dir / "profile"
        profile_dir.mkdir()
        shutil.copyfile(file_path, source_path)
        command = [
            libreoffice,
            "--headless",
            "--nologo",
            "--nodefault",
            "--nolockcheck",
            "--norestore",
            f"-env:UserInstallation={profile_dir.as_uri()}",
            "--convert-to",
            "xlsx",
            "--outdir",
            str(work_dir),
            str(source_path),
        ]
        try:
            result = subprocess.run(
                command,
                capture_output=True,
                text=True,
                timeout=CONVERSION_TIMEOUT_SECONDS,
                check=False,
            )
        except subprocess.TimeoutExpired as exc:
            raise SpreadsheetConversionError(
                f"XLS conversion timed out after {CONVERSION_TIMEOUT_SECONDS}s"
            ) from exc
        if result.returncode != 0 or not output_path.is_file():
            detail = result.stderr.strip() or result.stdout.strip() or "no XLSX output produced"
            raise SpreadsheetConversionError(f"XLS conversion failed: {detail[:500]}")
        return _parse_xlsx(str(output_path))


def parse_spreadsheet(file_path: str) -> Tuple[str, int]:
    ext = Path(file_path).suffix.lower()
    if ext == ".xlsx":
        text = _parse_xlsx(file_path)
    elif ext == ".xls":
        try:
            text = _parse_xls(file_path)
        except xlrd.biffh.XLRDError as exc:
            if "encrypted" not in str(exc).lower():
                raise
            logger.info(f"Using LibreOffice fallback for compatibility-encrypted XLS: {file_path}")
            text = _parse_compatibility_encrypted_xls(file_path)
    else:
        raise ValueError(f"unsupported spreadsheet type: {ext}")
    file_size = Path(file_path).stat().st_size
    logger.info(f"Parsed spreadsheet: {file_path}, text_len={len(text)}, size={file_size} bytes")
    return text, file_size
