"""Structured parser for PowerPoint Open XML presentations."""

from pathlib import Path
from typing import Tuple

from loguru import logger
from pptx import Presentation


MAX_EXTRACTED_CHARS = 10_000_000


def parse_presentation(file_path: str) -> Tuple[str, int]:
    presentation = Presentation(file_path)
    slides = []
    extracted_chars = 0
    for slide_number, slide in enumerate(presentation.slides, start=1):
        lines = [f"--- Slide {slide_number} ---"]
        for shape in slide.shapes:
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    while cells and not cells[-1]:
                        cells.pop()
                    if cells and any(cells):
                        lines.append("\t".join(cells))
                continue
            if getattr(shape, "has_text_frame", False):
                value = shape.text.strip()
                if value:
                    lines.append(value)
        rendered = "\n".join(lines)
        extracted_chars += len(rendered)
        if extracted_chars > MAX_EXTRACTED_CHARS:
            raise ValueError("presentation exceeds extracted text limit")
        slides.append(rendered)
    text = "\n\n".join(slides)
    file_size = Path(file_path).stat().st_size
    logger.info(f"Parsed presentation: {file_path}, slides={len(slides)}, size={file_size} bytes")
    return text, file_size
