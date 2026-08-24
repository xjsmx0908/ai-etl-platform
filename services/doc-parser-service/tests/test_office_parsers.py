import os

from openpyxl import Workbook
from pptx import Presentation
import xlwt

from app.routers.parse import ALLOWED_EXTENSIONS
from app.services.parser import parse_document
from app.services.parsers import spreadsheet


def test_parse_xlsx_preserves_sheet_and_table_structure(tmp_path):
    path = tmp_path / "policy.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "差旅标准"
    sheet.append(["城市级别", "住宿上限"])
    sheet.append(["一线城市", 600])
    second = workbook.create_sheet("审批规则")
    second.append(["金额", "审批人"])
    second.append([2000, "部门负责人"])
    workbook.save(path)

    text, file_size, parser_name = parse_document(str(path))

    assert "--- Sheet: 差旅标准 ---" in text
    assert "城市级别\t住宿上限" in text
    assert "一线城市\t600" in text
    assert "--- Sheet: 审批规则 ---" in text
    assert "2000\t部门负责人" in text
    assert file_size == path.stat().st_size
    assert parser_name == "xlsx"


def test_parse_legacy_xls_preserves_sheet_and_cells(tmp_path):
    path = tmp_path / "accounts.xls"
    workbook = xlwt.Workbook()
    sheet = workbook.add_sheet("账号清单")
    sheet.write(0, 0, "姓名")
    sheet.write(0, 1, "状态")
    sheet.write(1, 0, "测试用户")
    sheet.write(1, 1, "启用")
    workbook.save(str(path))

    text, file_size, parser_name = parse_document(str(path))

    assert "--- Sheet: 账号清单 ---" in text
    assert "姓名\t状态" in text
    assert "测试用户\t启用" in text
    assert file_size == path.stat().st_size
    assert parser_name == "xls"


def test_parse_compatibility_encrypted_xls_uses_libreoffice_fallback(tmp_path, monkeypatch):
    converted_xlsx = tmp_path / "converted.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "扣除信息"
    sheet.append(["项目", "年度"])
    sheet.append(["住房租金", 2026])
    workbook.save(converted_xlsx)

    converter = tmp_path / "libreoffice"
    converter.write_text(
        """#!/bin/sh
set -eu
outdir=""
while [ "$#" -gt 0 ]; do
    if [ "$1" = "--outdir" ]; then
        outdir="$2"
        shift 2
        continue
    fi
    shift
done
cp "$FAKE_XLSX" "$outdir/source.xlsx"
""",
        encoding="utf-8",
    )
    converter.chmod(0o755)
    monkeypatch.setenv("PATH", f"{tmp_path}:{os.environ['PATH']}")
    monkeypatch.setenv("FAKE_XLSX", str(converted_xlsx))
    monkeypatch.setattr(
        spreadsheet.xlrd,
        "open_workbook",
        lambda *args, **kwargs: (_ for _ in ()).throw(spreadsheet.xlrd.biffh.XLRDError("Workbook is encrypted")),
    )

    encrypted_xls = tmp_path / "encrypted.xls"
    encrypted_xls.write_bytes(bytes.fromhex("d0cf11e0a1b11ae1") + b"encrypted-probe")

    text, file_size, parser_name = parse_document(str(encrypted_xls))

    assert "--- Sheet: 扣除信息 ---" in text
    assert "住房租金\t2026" in text
    assert file_size == encrypted_xls.stat().st_size
    assert parser_name == "xls"


def test_parse_pptx_preserves_slide_text_and_tables(tmp_path):
    path = tmp_path / "training.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "财务制度培训"
    table = slide.shapes.add_table(2, 2, 0, 0, 4000000, 1000000).table
    table.cell(0, 0).text = "事项"
    table.cell(0, 1).text = "要求"
    table.cell(1, 0).text = "报销"
    table.cell(1, 1).text = "提交发票"
    presentation.save(path)

    text, file_size, parser_name = parse_document(str(path))

    assert "--- Slide 1 ---" in text
    assert "财务制度培训" in text
    assert "事项\t要求" in text
    assert "报销\t提交发票" in text
    assert file_size == path.stat().st_size
    assert parser_name == "pptx"


def test_office_extensions_are_accepted_at_parser_boundary():
    assert {".xls", ".xlsx", ".pptx"} <= ALLOWED_EXTENSIONS
